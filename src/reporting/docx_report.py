"""Professional Network Security Assessment DOCX Report Generator for ShieldScan.
Generates an executive-ready, highly polished Word report strictly honoring engineering design principles.
"""

from __future__ import annotations
from datetime import datetime
from pathlib import Path
from typing import Optional

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from src.models.assessment import ScanResult
from src.models.finding import FindingSeverity

# Color Palette for DOCX
HEX_TEXT_PRIMARY   = "1D1D1F"
HEX_TEXT_SECONDARY = "6E6E73"
HEX_ACCENT_RED     = "B42318"
HEX_BG_HEADER      = "F0F0EC"
HEX_BG_ALT         = "F9F9F8"
HEX_BORDER         = "D2D2CC"
HEX_SEV_HIGH       = "B42318"
HEX_SEV_MED        = "9A6700"
HEX_SEV_LOW        = "6E6E73"
HEX_SEV_INFO       = "2F6B4F"


def _set_cell_background(cell, hex_color: str):
    """Sets background color of a table cell via XML shading."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def _format_run(run, font_name="Calibri", font_size=10, bold=False, italic=False, color_rgb=(29, 29, 31)):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)


def generate_docx_report(scan_result: ScanResult, output_path: Path) -> Path:
    """Generates a complete, professional Network Security Assessment Word report."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # Page Margins
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    # ─────────────────────────────────────────────────────────────────────────
    # 01 — COVER PAGE
    # ─────────────────────────────────────────────────────────────────────────
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(72)
    p_pre.paragraph_format.space_after = Pt(12)
    r_sub = p_pre.add_run("SHIELDSCAN PLATFORM ASSESSMENT REPORT")
    _format_run(r_sub, font_size=11, bold=True, color_rgb=(180, 35, 24))

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run("Network Security & Visibility Assessment")
    _format_run(r_title, font_size=26, bold=True, color_rgb=(29, 29, 31))

    p_target = doc.add_paragraph()
    p_target.paragraph_format.space_after = Pt(40)
    r_tgt = p_target.add_run(f"Target Environment: {scan_result.target}")
    _format_run(r_tgt, font_size=13, color_rgb=(110, 110, 115))

    # Meta Table on Cover
    cover_table = doc.add_table(rows=5, cols=2)
    cover_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta_rows = [
        ("Assessment ID", scan_result.id),
        ("Execution Profile", scan_result.profile_name),
        ("Date & Time", scan_result.started_at[:19].replace("T", " ")),
        ("Total Duration", f"{scan_result.duration_seconds:.2f} seconds"),
        ("Assessed Hosts", f"{scan_result.stats.reachable_hosts} alive ({scan_result.stats.total_hosts} designated)"),
    ]
    for idx, (lbl, val) in enumerate(meta_rows):
        row = cover_table.rows[idx]
        _set_cell_margins(row.cells[0], top=80, bottom=80, left=100, right=100)
        _set_cell_margins(row.cells[1], top=80, bottom=80, left=100, right=100)
        _set_cell_background(row.cells[0], "F5F5F2")
        _set_cell_background(row.cells[1], "FFFFFF")
        
        p0 = row.cells[0].paragraphs[0]
        _format_run(p0.add_run(lbl), font_size=10, bold=True, color_rgb=(110, 110, 115))
        
        p1 = row.cells[1].paragraphs[0]
        _format_run(p1.add_run(val), font_size=10, color_rgb=(29, 29, 31))

    p_auth = doc.add_paragraph()
    p_auth.paragraph_format.space_before = Pt(80)
    r_auth = p_auth.add_run("Confidential Network Assessment · Authorized College Laboratory Demonstration")
    _format_run(r_auth, font_size=9, italic=True, color_rgb=(110, 110, 115))

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────────────────
    # 02 — EXECUTIVE SUMMARY
    # ─────────────────────────────────────────────────────────────────────────
    h1 = doc.add_heading(level=1)
    _format_run(h1.add_run("1. Executive Summary"), font_size=18, bold=True, color_rgb=(29, 29, 31))
    
    p_exec = doc.add_paragraph()
    p_exec.paragraph_format.space_after = Pt(14)
    r_ex_text = p_exec.add_run(
        f"This document presents the technical findings and network visibility intelligence collected by ShieldScan "
        f"during the assessment of {scan_result.target}. The evaluation was initiated under the '{scan_result.profile_name}' "
        f"profile, completing in {scan_result.duration_seconds:.2f} seconds. A total of {scan_result.stats.total_ports_tested:,} port "
        f"probes were conducted across {scan_result.stats.total_hosts} target hosts, uncovering {scan_result.stats.open_ports} open ports, "
        f"{scan_result.stats.services_identified} distinct active network services, and generating {scan_result.stats.total_findings} "
        f"evidence-grounded security findings."
    )
    _format_run(r_ex_text, font_size=10, color_rgb=(29, 29, 31))

    # Metric Snapshot Table
    sum_table = doc.add_table(rows=6, cols=3)
    sum_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Metric Category", "Observed Count", "Operational Context"]
    for c_idx, h_text in enumerate(headers):
        cell = sum_table.rows[0].cells[c_idx]
        _set_cell_background(cell, HEX_BG_HEADER)
        _set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        _format_run(p.add_run(h_text), font_size=10, bold=True, color_rgb=(29, 29, 31))

    sum_data = [
        ("Target Scope", f"{scan_result.stats.total_hosts} Hosts", f"{scan_result.stats.reachable_hosts} responsive, {scan_result.stats.total_hosts - scan_result.stats.reachable_hosts} unresponsive"),
        ("Port Exposure", f"{scan_result.stats.open_ports} Open Ports", f"{scan_result.stats.filtered_ports} filtered, {scan_result.stats.closed_ports} closed"),
        ("Services Identified", f"{scan_result.stats.services_identified} Services", "Fingerprinted via protocol banners and TLS"),
        ("Security Findings", f"{scan_result.stats.total_findings} Total", f"High: {scan_result.stats.findings_by_severity.get('HIGH',0)} | Med: {scan_result.stats.findings_by_severity.get('MEDIUM',0)} | Low: {scan_result.stats.findings_by_severity.get('LOW',0)}"),
        ("Informational Items", f"{scan_result.stats.findings_by_severity.get('INFORMATIONAL',0)} Items", "Verified service baselines and configurations"),
    ]

    for r_idx, row_data in enumerate(sum_data, start=1):
        row = sum_table.rows[r_idx]
        bg = HEX_BG_ALT if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            _set_cell_background(cell, bg)
            _set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            _format_run(p.add_run(str(val)), font_size=9.5, bold=(c_idx == 0), color_rgb=(29, 29, 31))

    # ─────────────────────────────────────────────────────────────────────────
    # 03 — SCOPE & METHODOLOGY
    # ─────────────────────────────────────────────────────────────────────────
    doc.add_paragraph().paragraph_format.space_before = Pt(16)
    h2 = doc.add_heading(level=1)
    _format_run(h2.add_run("2. Assessment Scope & Methodology"), font_size=18, bold=True, color_rgb=(29, 29, 31))

    p_meth = doc.add_paragraph()
    p_meth.paragraph_format.space_after = Pt(10)
    _format_run(p_meth.add_run(
        "ShieldScan employs a disciplined, non-destructive four-phase network visibility methodology "
        "designed for defensive verification:"
    ), font_size=10, color_rgb=(29, 29, 31))

    phases = [
        ("Phase 1: Multi-Vector Host Discovery", "Verifies host reachability via ICMP Echo requests, TCP SYN/ACK connects on administrative ports, and local reverse DNS resolution."),
        ("Phase 2: Asynchronous Port Enumeration", "Conducts non-blocking, managed-concurrency TCP probes across designated port lists, tracking latency and response states (Open, Closed, Filtered)."),
        ("Phase 3: Deep Service Fingerprinting", "Interrogates listening ports using standard protocol handshakes (HTTP HEAD, TLS certificate extraction, SSH banner analysis, MySQL handshakes, Redis commands), capturing raw evidence."),
        ("Phase 4: Evidence-First Security Analysis", "Translates concrete observations into risk classifications (Informational, Low, Medium, High) and correlates software versions against known CVE advisories without exploitation."),
    ]
    for ph_title, ph_desc in phases:
        p_ph = doc.add_paragraph(style='List Bullet')
        p_ph.paragraph_format.space_after = Pt(4)
        _format_run(p_ph.add_run(f"{ph_title}: "), font_size=10, bold=True, color_rgb=(29, 29, 31))
        _format_run(p_ph.add_run(ph_desc), font_size=9.5, color_rgb=(110, 110, 115))

    # ─────────────────────────────────────────────────────────────────────────
    # 04 — HOST INVENTORY
    # ─────────────────────────────────────────────────────────────────────────
    doc.add_paragraph().paragraph_format.space_before = Pt(16)
    h3 = doc.add_heading(level=1)
    _format_run(h3.add_run("3. Host Inventory"), font_size=18, bold=True, color_rgb=(29, 29, 31))

    host_table = doc.add_table(rows=len(scan_result.hosts) + 1, cols=6)
    host_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_headers = ["IP Address", "Hostname", "Status", "Latency", "Open Ports", "Findings"]
    for c_idx, h_text in enumerate(h_headers):
        cell = host_table.rows[0].cells[c_idx]
        _set_cell_background(cell, HEX_BG_HEADER)
        _set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        _format_run(p.add_run(h_text), font_size=9.5, bold=True, color_rgb=(29, 29, 31))

    for r_idx, h in enumerate(scan_result.hosts.values(), start=1):
        row = host_table.rows[r_idx]
        bg = HEX_BG_ALT if r_idx % 2 == 0 else "FFFFFF"
        row_vals = [
            h.ip,
            h.hostname or "—",
            h.status.value,
            f"{h.latency_ms:.1f} ms",
            str(len(h.open_ports)),
            str(len(h.findings)),
        ]
        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            _set_cell_background(cell, bg)
            _set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            _format_run(p.add_run(val), font_size=9, bold=(c_idx == 0), color_rgb=(29, 29, 31))

    # ─────────────────────────────────────────────────────────────────────────
    # 05 — PORT & SERVICE INVENTORY
    # ─────────────────────────────────────────────────────────────────────────
    doc.add_paragraph().paragraph_format.space_before = Pt(16)
    h4 = doc.add_heading(level=1)
    _format_run(h4.add_run("4. Port & Service Inventory"), font_size=18, bold=True, color_rgb=(29, 29, 31))

    svc_list = scan_result.all_services
    if not svc_list:
        p_nosvc = doc.add_paragraph()
        _format_run(p_nosvc.add_run("No active listening services were detected on assessed target hosts."), font_size=10, italic=True)
    else:
        svc_table = doc.add_table(rows=len(svc_list) + 1, cols=6)
        svc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        s_headers = ["Host", "Port / Proto", "Service", "Product & Version", "Confidence", "Latency"]
        for c_idx, h_text in enumerate(s_headers):
            cell = svc_table.rows[0].cells[c_idx]
            _set_cell_background(cell, HEX_BG_HEADER)
            _set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
            p = cell.paragraphs[0]
            _format_run(p.add_run(h_text), font_size=9.5, bold=True, color_rgb=(29, 29, 31))

        for r_idx, s in enumerate(svc_list, start=1):
            row = svc_table.rows[r_idx]
            bg = HEX_BG_ALT if r_idx % 2 == 0 else "FFFFFF"
            prod_ver = f"{s['product']} {s['version']}".strip() or "—"
            s_vals = [
                s["host"],
                f"{s['port']}/{s['protocol']}",
                s["name"],
                prod_ver,
                s["confidence"],
                f"{s['latency_ms']:.1f} ms",
            ]
            for c_idx, val in enumerate(s_vals):
                cell = row.cells[c_idx]
                _set_cell_background(cell, bg)
                _set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
                p = cell.paragraphs[0]
                _format_run(p.add_run(val), font_size=9, bold=(c_idx == 1), color_rgb=(29, 29, 31))

    # ─────────────────────────────────────────────────────────────────────────
    # 06 — SECURITY FINDINGS (EVIDENCE-GROUNDED)
    # ─────────────────────────────────────────────────────────────────────────
    doc.add_paragraph().paragraph_format.space_before = Pt(16)
    h5 = doc.add_heading(level=1)
    _format_run(h5.add_run("5. Security Findings & Evidence Trail"), font_size=18, bold=True, color_rgb=(29, 29, 31))

    findings = scan_result.all_findings
    if not findings:
        p_nofind = doc.add_paragraph()
        _format_run(p_nofind.add_run("No anomalous security conditions or exposed vulnerabilities detected."), font_size=10, italic=True)
    else:
        for f in findings:
            # Finding Card Table
            f_table = doc.add_table(rows=5, cols=2)
            f_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            # Severity color mapping
            sev_color = (180, 35, 24) if f.severity == FindingSeverity.HIGH else (
                (154, 103, 0) if f.severity == FindingSeverity.MEDIUM else (
                    (47, 107, 79) if f.severity == FindingSeverity.INFORMATIONAL else (110, 110, 115)
                )
            )

            # Header row
            hdr_cell = f_table.rows[0].cells[0]
            _set_cell_background(hdr_cell, "F0F0EC")
            _set_cell_margins(hdr_cell, top=100, bottom=100, left=120, right=120)
            p_h = hdr_cell.paragraphs[0]
            _format_run(p_h.add_run(f"[{f.severity.value}]  {f.id} — {f.title}"), font_size=11, bold=True, color_rgb=sev_color)

            # Target cell
            tgt_cell = f_table.rows[0].cells[1]
            _set_cell_background(tgt_cell, "F0F0EC")
            _set_cell_margins(tgt_cell, top=100, bottom=100, left=120, right=120)
            p_t = tgt_cell.paragraphs[0]
            p_t.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _format_run(p_t.add_run(f"{f.host}:{f.port if f.port else '—'} ({f.service or 'System'})"), font_size=9.5, color_rgb=(110, 110, 115))

            # Description
            d_lbl = f_table.rows[1].cells[0]
            _format_run(d_lbl.paragraphs[0].add_run("Description"), font_size=9, bold=True, color_rgb=(110, 110, 115))
            d_val = f_table.rows[1].cells[1]
            _format_run(d_val.paragraphs[0].add_run(f.description), font_size=9.5, color_rgb=(29, 29, 31))

            # Technical Evidence
            ev_lbl = f_table.rows[2].cells[0]
            _format_run(ev_lbl.paragraphs[0].add_run("Observed Evidence"), font_size=9, bold=True, color_rgb=(110, 110, 115))
            ev_val = f_table.rows[2].cells[1]
            ev_text = f.evidence.observation if f.evidence else "Observed during active port discovery"
            if f.evidence and f.evidence.raw_data:
                ev_text += f"\n[Raw Trace: {f.evidence.raw_data[:120]}]"
            _format_run(ev_val.paragraphs[0].add_run(ev_text), font_size=9, color_rgb=(29, 29, 31))

            # Impact
            imp_lbl = f_table.rows[3].cells[0]
            _format_run(imp_lbl.paragraphs[0].add_run("Potential Impact"), font_size=9, bold=True, color_rgb=(110, 110, 115))
            imp_val = f_table.rows[3].cells[1]
            _format_run(imp_val.paragraphs[0].add_run(f.impact or "Information disclosure or unauthorized service enumeration."), font_size=9.5, color_rgb=(29, 29, 31))

            # Remediation
            rec_lbl = f_table.rows[4].cells[0]
            _format_run(rec_lbl.paragraphs[0].add_run("Remediation"), font_size=9, bold=True, color_rgb=(47, 107, 79))
            rec_val = f_table.rows[4].cells[1]
            _format_run(rec_val.paragraphs[0].add_run(f.recommendation or "Audit service authorization and enforce access controls."), font_size=9.5, bold=True, color_rgb=(29, 29, 31))

            for r in f_table.rows:
                _set_cell_margins(r.cells[0], top=80, bottom=80, left=100, right=100)
                _set_cell_margins(r.cells[1], top=80, bottom=80, left=100, right=100)

            doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ─────────────────────────────────────────────────────────────────────────
    # 07 — STRATEGIC RECOMMENDATIONS
    # ─────────────────────────────────────────────────────────────────────────
    h6 = doc.add_heading(level=1)
    _format_run(h6.add_run("6. Strategic Remediation Roadmap"), font_size=18, bold=True, color_rgb=(29, 29, 31))

    p_recs_intro = doc.add_paragraph()
    _format_run(p_recs_intro.add_run(
        "Based on the specific services and risk exposures identified during this assessment, "
        "the following prioritized engineering countermeasures should be implemented:"
    ), font_size=10, color_rgb=(29, 29, 31))

    recs = [
        ("Enforce Modern Transport Encryption", "Deprecate plaintext transport protocols (such as unencrypted HTTP, Telnet, or FTP). Mandate TLS 1.3 or SSH v2 for all administrative and user communications."),
        ("Isolate Database Management Ports", "Ensure database listening ports (e.g., MySQL 3306, Redis 6379, PostgreSQL 5432) are bound exclusively to 127.0.0.1 or restricted private VLANs with strict host firewall filtering."),
        ("Bastion & VPN for Remote Administration", "Restrict administrative surfaces (RDP, SSH, WinRM) behind multi-factor authenticated VPN gateways rather than exposing them to wide network segments."),
        ("Continuous Discovery & Differential Audits", "Establish periodic ShieldScan automated assessments to track host and port drift, catching newly exposed listeners before they are discovered externally."),
    ]
    for r_title, r_desc in recs:
        p_r = doc.add_paragraph(style='List Bullet')
        p_r.paragraph_format.space_after = Pt(6)
        _format_run(p_r.add_run(f"{r_title}: "), font_size=10, bold=True, color_rgb=(29, 29, 31))
        _format_run(p_r.add_run(r_desc), font_size=9.5, color_rgb=(110, 110, 115))

    # ─────────────────────────────────────────────────────────────────────────
    # 08 — TECHNICAL APPENDIX & LIMITATIONS
    # ─────────────────────────────────────────────────────────────────────────
    doc.add_paragraph().paragraph_format.space_before = Pt(16)
    h7 = doc.add_heading(level=1)
    _format_run(h7.add_run("7. Technical Appendix & Assessment Limitations"), font_size=18, bold=True, color_rgb=(29, 29, 31))

    p_app = doc.add_paragraph()
    _format_run(p_app.add_run(
        f"Scan Engine: ShieldScan Enterprise Core v5.0\n"
        f"Assessment Profile: {scan_result.profile_name} (ID: {scan_result.profile_id})\n"
        f"Probe Timeout: {scan_result.config_summary.get('timeout', 0.35)}s | Concurrency: {scan_result.config_summary.get('concurrency', 250)} workers\n"
        f"Timestamp: Started at {scan_result.started_at} | Concluded at {scan_result.completed_at or 'N/A'}\n"
    ), font_size=9, color_rgb=(110, 110, 115))

    p_lim_title = doc.add_paragraph()
    _format_run(p_lim_title.add_run("Assessment Limitations:"), font_size=10, bold=True, color_rgb=(29, 29, 31))

    limits = [
        "Network Firewalls & Stateful Inspection: In-line firewalls or host-based packet filters (e.g., Windows Defender Firewall, iptables) may drop or silently reset probes, causing ports to be categorized as FILTERED.",
        "Fingerprint Uncertainty: Version identification depends on banners published by target daemons. Custom or obfuscated banners may produce LOW confidence classifications.",
        "Defensive Scope: ShieldScan operates strictly as a discovery, enumeration, and non-destructive assessment tool. It does not perform active exploits, denial of service, or authentication bypass.",
    ]
    for l_text in limits:
        p_l = doc.add_paragraph(style='List Bullet')
        p_l.paragraph_format.space_after = Pt(4)
        _format_run(p_l.add_run(l_text), font_size=9, color_rgb=(110, 110, 115))

    doc.save(str(output_path))
    return output_path

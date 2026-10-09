"""College Technical Presentation Generator for ShieldScan.
Generates an executive-grade 7-slide PPTX deck derived directly from canonical ScanResult data.
Strictly adheres to Apple-inspired minimalism and the ShieldScan color lock.
"""

from __future__ import annotations
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from src.models.assessment import ScanResult
from src.presentation.diagrams import generate_architecture_diagram, generate_pipeline_diagram

# Palette Tokens (Strict Lock)
C_CANVAS   = RGBColor(245, 245, 242)   # #F5F5F2
C_SURFACE  = RGBColor(255, 255, 255)   # #FFFFFF
C_TEXT     = RGBColor(29, 29, 31)      # #1D1D1F
C_MUTED    = RGBColor(110, 110, 115)   # #6E6E73
C_ACCENT   = RGBColor(180, 35, 24)     # #B42318
C_GREEN    = RGBColor(47, 107, 79)     # #2F6B4F
C_BORDER   = RGBColor(210, 210, 204)   # #D2D2CC


def _set_slide_bg(slide):
    """Fills slide background with warm Apple canvas (#F5F5F2)."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = C_CANVAS


def _add_card(slide, left, top, width, height, border_color=C_BORDER, fill_color=C_SURFACE):
    """Draws a clean rectangular surface card."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.0)
    return shape


def generate_presentation(scan_result: ScanResult, output_path: Path, screenshots_dir: Optional[Path] = None) -> Path:
    """Generates the comprehensive 7-slide academic presentation."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Pre-generate diagrams
    assets_dir = output_path.parent / "assets"
    arch_png = generate_architecture_diagram(assets_dir / "architecture.png")
    pipe_png = generate_pipeline_diagram(assets_dir / "pipeline.png")

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 1 — TITLE
    # ═════════════════════════════════════════════════════════════════════════
    s1 = prs.slides.add_slide(blank_layout)
    _set_slide_bg(s1)

    # Accent pill
    tb_pill = s1.shapes.add_textbox(Inches(1.2), Inches(1.4), Inches(6.0), Inches(0.5))
    p_pill = tb_pill.text_frame.paragraphs[0]
    r_pill = p_pill.add_run()
    r_pill.text = "COMPUTER NETWORKS PROJECT · ACADEMIC DEFENSE"
    r_pill.font.size = Pt(11)
    r_pill.font.bold = True
    r_pill.font.color.rgb = C_ACCENT

    # Main Title
    tb_title = s1.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(11.0), Inches(1.8))
    p_t = tb_title.text_frame.paragraphs[0]
    r_t = p_t.add_run()
    r_t.text = "SHIELDSCAN"
    r_t.font.size = Pt(56)
    r_t.font.bold = True
    r_t.font.color.rgb = C_TEXT

    p_sub = tb_title.text_frame.add_paragraph()
    r_sub = p_sub.add_run()
    r_sub.text = "Modern Network Visibility & Security Assessment Platform"
    r_sub.font.size = Pt(22)
    r_sub.font.color.rgb = C_MUTED

    # Pipeline graphic
    s1.shapes.add_picture(str(pipe_png), Inches(1.2), Inches(4.2), Inches(10.8), Inches(2.3))

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 2 — THE PROBLEM
    # ═════════════════════════════════════════════════════════════════════════
    s2 = prs.slides.add_slide(blank_layout)
    _set_slide_bg(s2)

    tb2 = s2.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.0), Inches(1.0))
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = "The Network Assessment Challenge"
    p2.font.size = Pt(30)
    p2.font.bold = True
    p2.font.color.rgb = C_TEXT

    p2_sub = tb2.text_frame.add_paragraph()
    p2_sub.text = "Traditional tools generate overwhelming raw output without context, evidence, or historical intelligence."
    p2_sub.font.size = Pt(14)
    p2_sub.font.color.rgb = C_MUTED

    # 3 Comparison Problem Cards
    cards_data = [
        ("Raw Scanner Output", "Terminal tools dump thousands of lines of raw text. Security operators must manually parse ports and look up service banners.", C_BORDER),
        ("Unverifiable Claims", "Conventional tools report speculative vulnerabilities without providing auditable evidence of how or why the finding was derived.", C_BORDER),
        ("Zero Historical Context", "Point-in-time scans do not track network drift. Identifying newly exposed ports or resolved risks requires tedious manual diffs.", C_ACCENT),
    ]
    for i, (ctitle, cdesc, border_c) in enumerate(cards_data):
        cx = Inches(1.0 + (i * 3.85))
        _add_card(s2, cx, Inches(2.2), Inches(3.6), Inches(4.3), border_color=border_c)
        tb_c = s2.shapes.add_textbox(cx + Inches(0.3), Inches(2.5), Inches(3.0), Inches(3.6))
        
        p_c1 = tb_c.text_frame.paragraphs[0]
        p_c1.text = f"0{i+1}"
        p_c1.font.size = Pt(20)
        p_c1.font.bold = True
        p_c1.font.color.rgb = C_ACCENT if border_c == C_ACCENT else C_MUTED
        
        p_c2 = tb_c.text_frame.add_paragraph()
        p_c2.space_before = Pt(14)
        p_c2.text = ctitle
        p_c2.font.size = Pt(16)
        p_c2.font.bold = True
        p_c2.font.color.rgb = C_TEXT
        
        p_c3 = tb_c.text_frame.add_paragraph()
        p_c3.space_before = Pt(10)
        p_c3.text = cdesc
        p_c3.font.size = Pt(12)
        p_c3.font.color.rgb = C_MUTED

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 3 — OBJECTIVES
    # ═════════════════════════════════════════════════════════════════════════
    s3 = prs.slides.add_slide(blank_layout)
    _set_slide_bg(s3)

    tb3 = s3.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.0), Inches(1.0))
    p3 = tb3.text_frame.paragraphs[0]
    p3.text = "System Objectives & Engineering Scope"
    p3.font.size = Pt(30)
    p3.font.bold = True
    p3.font.color.rgb = C_TEXT

    p3_sub = tb3.text_frame.add_paragraph()
    p3_sub.text = "Designing a non-destructive, evidence-driven network assessment architecture."
    p3_sub.font.size = Pt(14)
    p3_sub.font.color.rgb = C_MUTED

    objs = [
        ("Multi-Vector Discovery", "Verify alive systems via ICMP Echo, TCP SYN connects, and reverse DNS across subnets without hanging."),
        ("Asynchronous Enumeration", "Conduct high-speed port scanning with controlled concurrency pools and adaptive timeout handling."),
        ("Deep Service Fingerprinting", "Interrogate services using live protocol handshakes (HTTP, TLS certificates, SSH banners, MySQL)."),
        ("Evidence-Based Findings", "Classify exposures strictly supported by observed data into Informational, Low, Medium, and High risk."),
        ("Historical Drift Detection", "Track differential changes (+/- hosts, +/- open ports, version upgrades) across consecutive assessments."),
        ("Automated Executive Reporting", "Generate single-source-of-truth DOCX reports and structured exports instantly with zero placeholders."),
    ]
    for idx, (otitle, odesc) in enumerate(objs):
        col = idx % 2
        row = idx // 2
        ox = Inches(1.0 + (col * 5.8))
        oy = Inches(2.2 + (row * 1.5))
        _add_card(s3, ox, oy, Inches(5.5), Inches(1.3))
        tb_o = s3.shapes.add_textbox(ox + Inches(0.2), oy + Inches(0.15), Inches(5.1), Inches(1.0))
        po1 = tb_o.text_frame.paragraphs[0]
        po1.text = otitle
        po1.font.size = Pt(14)
        po1.font.bold = True
        po1.font.color.rgb = C_TEXT
        po2 = tb_o.text_frame.add_paragraph()
        po2.text = odesc
        po2.font.size = Pt(10.5)
        po2.font.color.rgb = C_MUTED

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 4 — ARCHITECTURE
    # ═════════════════════════════════════════════════════════════════════════
    s4 = prs.slides.add_slide(blank_layout)
    _set_slide_bg(s4)

    tb4 = s4.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.0), Inches(0.8))
    p4 = tb4.text_frame.paragraphs[0]
    p4.text = "ShieldScan Product Architecture"
    p4.font.size = Pt(28)
    p4.font.bold = True
    p4.font.color.rgb = C_TEXT

    # Embed architecture diagram
    s4.shapes.add_picture(str(arch_png), Inches(1.0), Inches(1.6), Inches(11.3), Inches(5.3))

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 5 — REAL SCAN ASSESSMENT & METRICS
    # ═════════════════════════════════════════════════════════════════════════
    s5 = prs.slides.add_slide(blank_layout)
    _set_slide_bg(s5)

    tb5 = s5.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.0), Inches(1.0))
    p5 = tb5.text_frame.paragraphs[0]
    p5.text = f"Live Network Assessment: {scan_result.target}"
    p5.font.size = Pt(30)
    p5.font.bold = True
    p5.font.color.rgb = C_TEXT

    p5_sub = tb5.text_frame.add_paragraph()
    p5_sub.text = f"Assessment completed in {scan_result.duration_seconds:.2f}s under the '{scan_result.profile_name}' profile."
    p5_sub.font.size = Pt(14)
    p5_sub.font.color.rgb = C_MUTED

    # 4 Real Metric Big Number Cards
    metrics = [
        (str(scan_result.stats.reachable_hosts), "Hosts Discovered", f"Out of {scan_result.stats.total_hosts} scope targets"),
        (str(scan_result.stats.open_ports), "Open Ports", f"{scan_result.stats.filtered_ports} filtered / firewalled"),
        (str(scan_result.stats.services_identified), "Services Identified", "Fingerprinted via banners & TLS"),
        (str(scan_result.stats.total_findings), "Security Findings", f"High: {scan_result.stats.findings_by_severity.get('HIGH',0)} | Med: {scan_result.stats.findings_by_severity.get('MEDIUM',0)}"),
    ]
    for i, (mval, mlbl, mctx) in enumerate(metrics):
        mx = Inches(1.0 + (i * 2.9))
        _add_card(s5, mx, Inches(2.2), Inches(2.65), Inches(2.2))
        tb_m = s5.shapes.add_textbox(mx + Inches(0.2), Inches(2.4), Inches(2.25), Inches(1.8))
        
        pm1 = tb_m.text_frame.paragraphs[0]
        pm1.text = mval
        pm1.font.size = Pt(38)
        pm1.font.bold = True
        pm1.font.color.rgb = C_ACCENT if i == 3 and int(mval) > 0 else C_TEXT
        
        pm2 = tb_m.text_frame.add_paragraph()
        pm2.text = mlbl
        pm2.font.size = Pt(13)
        pm2.font.bold = True
        pm2.font.color.rgb = C_TEXT
        
        pm3 = tb_m.text_frame.add_paragraph()
        pm3.text = mctx
        pm3.font.size = Pt(9.5)
        pm3.font.color.rgb = C_MUTED

    # Check if dashboard screenshot exists
    default_ss_dir = Path(__file__).resolve().parents[2] / "screenshots"
    dash_ss = (screenshots_dir or default_ss_dir) / "01_dashboard_overview.png"
    if dash_ss.exists():

        # Right side: Screenshot card
        s5.shapes.add_picture(str(dash_ss), Inches(6.8), Inches(4.7), Inches(5.53), Inches(2.2))
        # Left side: Assessment Highlights Card
        _add_card(s5, Inches(1.0), Inches(4.7), Inches(5.6), Inches(2.2))
        tb_bot = s5.shapes.add_textbox(Inches(1.2), Inches(4.8), Inches(5.2), Inches(2.0))
    else:
        # Full width Assessment Highlights Card
        _add_card(s5, Inches(1.0), Inches(4.7), Inches(11.33), Inches(2.2))
        tb_bot = s5.shapes.add_textbox(Inches(1.3), Inches(4.9), Inches(10.7), Inches(1.8))
    pb1 = tb_bot.text_frame.paragraphs[0]
    pb1.text = "Key Observational Intelligence"
    pb1.font.size = Pt(15)
    pb1.font.bold = True
    pb1.font.color.rgb = C_TEXT

    # List top findings or services
    sample_findings = scan_result.all_findings[:3]
    if sample_findings:
        for f in sample_findings:
            pbf = tb_bot.text_frame.add_paragraph()
            pbf.text = f"• [{f.severity.value}] {f.title} ({f.host}:{f.port or 'System'}) — Evidence: {f.evidence.observation[:90] if f.evidence else 'Verified'}"
            pbf.font.size = Pt(11)
            pbf.font.color.rgb = C_TEXT
    else:
        pbf = tb_bot.text_frame.add_paragraph()
        pbf.text = f"• All {scan_result.stats.open_ports} active listening ports verified with non-destructive protocol probes."
        pbf.font.size = Pt(11)
        pbf.font.color.rgb = C_TEXT

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 6 — DIFFERENTIATION: WHY SHIELDSCAN?
    # ═════════════════════════════════════════════════════════════════════════
    s6 = prs.slides.add_slide(blank_layout)
    _set_slide_bg(s6)

    tb6 = s6.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.0), Inches(1.0))
    p6 = tb6.text_frame.paragraphs[0]
    p6.text = "Why ShieldScan? Complete Network Intelligence"
    p6.font.size = Pt(30)
    p6.font.bold = True
    p6.font.color.rgb = C_TEXT

    p6_sub = tb6.text_frame.add_paragraph()
    p6_sub.text = "SCAN + VISUALIZE + ANALYZE + COMPARE + REPORT in one unified product experience."
    p6_sub.font.size = Pt(14)
    p6_sub.font.color.rgb = C_MUTED

    difs = [
        ("Evidence-First Integrity", "Never manufactures findings or CVE matches. Every reported risk is traceable directly to an observed raw banner, TLS handshake, or protocol packet."),
        ("Integrated Visual Dashboard", "Replaces cluttered CLI flags and hacker aesthetics with an Apple-inspired clean UI designed for clarity, progressive depth, and instant comprehension."),
        ("Automated Change Detection", "Instantly diffs consecutive network scans to highlight newly opened ports, removed hosts, and version migrations across time."),
        ("One-Click Executive Reporting", "Derives professional Word (.docx) reports and slide presentations directly from canonical scan data with zero manual transposition."),
    ]
    for i, (dtitle, ddesc) in enumerate(difs):
        dx = Inches(1.0 + (i % 2 * 5.8))
        dy = Inches(2.2 + (i // 2 * 2.3))
        _add_card(s6, dx, dy, Inches(5.5), Inches(2.0))
        tb_d = s6.shapes.add_textbox(dx + Inches(0.3), dy + Inches(0.2), Inches(4.9), Inches(1.6))
        pd1 = tb_d.text_frame.paragraphs[0]
        pd1.text = dtitle
        pd1.font.size = Pt(15)
        pd1.font.bold = True
        pd1.font.color.rgb = C_ACCENT if i == 0 else C_TEXT
        pd2 = tb_d.text_frame.add_paragraph()
        pd2.space_before = Pt(8)
        pd2.text = ddesc
        pd2.font.size = Pt(11)
        pd2.font.color.rgb = C_MUTED

    # ═════════════════════════════════════════════════════════════════════════
    # SLIDE 7 — CONCLUSION & FUTURE SCOPE
    # ═════════════════════════════════════════════════════════════════════════
    s7 = prs.slides.add_slide(blank_layout)
    _set_slide_bg(s7)

    tb7 = s7.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.0), Inches(1.0))
    p7 = tb7.text_frame.paragraphs[0]
    p7.text = "From Raw Network Data to Understanding"
    p7.font.size = Pt(30)
    p7.font.bold = True
    p7.font.color.rgb = C_TEXT

    p7_sub = tb7.text_frame.add_paragraph()
    p7_sub.text = "Empowering defensive engineers with accessible network clarity and reliable evidence."
    p7_sub.font.size = Pt(14)
    p7_sub.font.color.rgb = C_MUTED

    # Summary card left
    _add_card(s7, Inches(1.0), Inches(2.2), Inches(5.5), Inches(4.3))
    tb_sum = s7.shapes.add_textbox(Inches(1.3), Inches(2.5), Inches(4.9), Inches(3.7))
    ps1 = tb_sum.text_frame.paragraphs[0]
    ps1.text = "Project Summary"
    ps1.font.size = Pt(18)
    ps1.font.bold = True
    ps1.font.color.rgb = C_TEXT

    points = [
        "Transformed port scanning into a comprehensive network visibility platform.",
        "Strictly enforced Apple-inspired light engineering UI (zero dark/neon hacker styling).",
        "Established single-source-of-truth canonical ScanResult model for all reports and UI.",
        "Demonstrated defensive network assessment without offensive exploitation.",
    ]
    for pt in points:
        psp = tb_sum.text_frame.add_paragraph()
        psp.space_before = Pt(8)
        psp.text = f"✓  {pt}"
        psp.font.size = Pt(11.5)
        psp.font.color.rgb = C_TEXT

    # Future scope card right
    _add_card(s7, Inches(6.8), Inches(2.2), Inches(5.5), Inches(4.3), border_color=C_ACCENT)
    tb_fut = s7.shapes.add_textbox(Inches(7.1), Inches(2.5), Inches(4.9), Inches(3.7))
    pf1 = tb_fut.text_frame.paragraphs[0]
    pf1.text = "Future Roadmap"
    pf1.font.size = Pt(18)
    pf1.font.bold = True
    pf1.font.color.rgb = C_ACCENT

    fut_points = [
        "Heuristic OS Fingerprinting: TCP/IP stack window size and TTL analysis.",
        "Interactive Topology Graph: Force-directed visual mapping of subnet relationships.",
        "Continuous Distributed Agent: Background daemon monitoring for enterprise networks.",
        "Expanded Protocol Probing: Native SMB2/3 tree connect and LDAP schema enumeration.",
    ]
    for fp in fut_points:
        pfp = tb_fut.text_frame.add_paragraph()
        pfp.space_before = Pt(8)
        pfp.text = f"→  {fp}"
        pfp.font.size = Pt(11.5)
        pfp.font.color.rgb = C_MUTED

    prs.save(str(output_path))
    return output_path
